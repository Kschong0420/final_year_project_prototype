"""Milestone 3 HTTP/WebSocket integration checks with generated, disposable files."""
import io
import json
import os
import sys
import unittest
from contextlib import ExitStack
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pymupdf
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app.main import create_app


def pdf_bytes():
    document = pymupdf.open()
    for text in ("First PDF page", "Second PDF page"):
        page = document.new_page()
        page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def pptx_bytes():
    presentation = Presentation()
    for text in ("First PPTX slide", "Second PPTX slide"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text = text
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def receive(socket, predicate=lambda state: True):
    for _ in range(30):
        state = socket.receive_json()
        if predicate(state):
            return state
    raise AssertionError("Expected WebSocket state was not received")


class MaterialTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(create_app())
        self.lecturer = self.client.post("/api/sessions", json={"title": "Documents"}).json()
        self.code = self.lecturer["code"]
        self.upload_url = f"/api/sessions/{self.code}/materials"
        self.headers = {"Authorization": f"Bearer {self.lecturer['token']}"}

    def tearDown(self):
        self.client.close()

    def upload(self, name, content, headers=None):
        return self.client.post(self.upload_url, headers=self.headers if headers is None else headers,
                                files={"file": (name, content, "application/octet-stream")})

    def join(self):
        return self.client.post(f"/api/sessions/{self.code}/join").json()

    def connect(self, stack, credentials):
        socket = stack.enter_context(self.client.websocket_connect(f"/ws/sessions/{self.code}"))
        socket.send_json({"token": credentials["token"]})
        self.last_connected_state = receive(socket, lambda state: state.get("type") == "state")
        return socket

    def test_pdf_upload_sync_feedback_replacement_and_access(self):
        with ExitStack() as stack:
            teacher = self.connect(stack, self.lecturer)
            student_a = self.join()
            student_b = self.join()
            a = self.connect(stack, student_a)
            b = self.connect(stack, student_b)
            initial = receive(teacher, lambda s: s.get("connected_students") == 2)
            self.assertEqual(len(initial["slides"]), 5)
            result = self.upload("lecture.pdf", pdf_bytes())
            self.assertEqual(result.status_code, 201, result.text)
            self.assertEqual(result.json()["slide_count"], 2)
            current = receive(teacher, lambda s: s.get("presentation_id") == result.json()["presentation_id"])
            self.assertEqual(current["current_slide"], 0)
            self.assertEqual(len(current["slides"]), 2)
            self.assertIn("First PDF page", current["slides"][0]["text"])
            self.assertIn("Second PDF page", current["slides"][1]["text"])
            self.assertEqual(current["current_feedback"]["total"], 0)
            image_url = current["slides"][0]["image_url"]
            self.assertEqual(self.client.get(image_url).status_code, 403)
            self.assertEqual(self.client.get(image_url, headers={"Authorization": f"Bearer {student_a['token']}"}).status_code, 200)
            self.assertEqual(self.client.get(image_url, headers=self.headers).headers["content-type"], "image/png")
            receive(a, lambda s: s.get("presentation_id") == current["presentation_id"])
            receive(b, lambda s: s.get("presentation_id") == current["presentation_id"])

            teacher.send_json({"type": "set_slide", "index": 1})
            for socket in (teacher, a, b):
                self.assertEqual(receive(socket, lambda s: s.get("current_slide") == 1)["current_slide"], 1)
            a.send_json({"type": "set_slide", "index": 0})
            self.assertIn("lecturer", receive(a, lambda s: s.get("type") == "error")["message"])
            late_credentials = self.join()
            late = self.connect(stack, late_credentials)
            self.assertEqual(self.last_connected_state["current_slide"], 1)
            a.close()
            reconnected = self.connect(stack, student_a)
            self.assertEqual(self.last_connected_state["current_slide"], 1)

            for socket, choice in ((reconnected, "not_understand"), (b, "understand")):
                socket.send_json({"type": "submit_feedback", "slide_index": 1, "choice": choice,
                                  "presentation_id": current["presentation_id"]})
            flagged = receive(teacher, lambda s: s.get("current_feedback", {}).get("total") == 2)
            self.assertTrue(flagged["current_feedback"]["flagged"])
            teacher.send_json({"type": "set_slide", "index": 0})
            empty = receive(teacher, lambda s: s.get("current_slide") == 0)
            self.assertEqual(empty["current_feedback"]["total"], 0)
            teacher.send_json({"type": "set_slide", "index": 1})
            self.assertEqual(receive(teacher, lambda s: s.get("current_slide") == 1)["current_feedback"]["total"], 2)

            another = self.upload("second.pdf", pdf_bytes())
            self.assertEqual(another.status_code, 201, another.text)
            replacement = receive(teacher, lambda s: s.get("presentation_id") == another.json()["presentation_id"])
            self.assertEqual(replacement["current_slide"], 0)
            self.assertEqual(replacement["current_feedback"]["total"], 0)
            self.assertEqual(replacement["flagged_slides"], [])
            self.assertEqual(self.client.get(image_url, headers=self.headers).status_code, 404)
            reconnected.send_json({"type": "submit_feedback", "slide_index": 0, "choice": "understand",
                                   "presentation_id": current["presentation_id"]})
            self.assertIn("presentation has changed", receive(reconnected, lambda s: s.get("type") == "error")["message"])
            teacher.send_json({"type": "set_slide", "index": 1})
            self.assertEqual(receive(teacher, lambda s: s.get("presentation_id") == replacement["presentation_id"] and s.get("current_slide") == 1)["current_feedback"]["total"], 0)

    def test_pptx_order_and_failure_keeps_previous(self):
        with patch("app.document_processing.libreoffice_executable", return_value=None):
            good = self.upload("lecture.pptx", pptx_bytes())
        self.assertEqual(good.status_code, 201, good.text)
        self.assertEqual(good.json()["rendering"], "text-only")
        self.assertIn("LibreOffice", good.json()["warning"])
        with ExitStack() as stack:
            teacher = self.connect(stack, self.lecturer)
            student_a = self.connect(stack, self.join())
            student_b = self.connect(stack, self.join())
            teacher.send_json({"type": "ping"})
            receive(teacher, lambda s: s.get("type") == "pong")
            state = self.client.app.state.store.session.snapshot("lecturer", self.lecturer["token"])
            self.assertEqual(len(state["slides"]), 2)
            self.assertEqual(state["slides"][0]["points"], ["First PPTX slide"])
            self.assertEqual(state["slides"][1]["points"], ["Second PPTX slide"])
            self.assertIsNone(state["slides"][0]["image_url"])
            teacher.send_json({"type": "set_slide", "index": 1})
            state = receive(teacher, lambda s: s.get("current_slide") == 1)
            for student in (student_a, student_b):
                self.assertEqual(receive(student, lambda s: s.get("current_slide") == 1)["presentation_id"],
                                 state["presentation_id"])
            for name, content, status in (("old.ppt", b"old", 415), ("bad.txt", b"bad", 415),
                                           ("bad.pdf", b"broken", 422), ("bad.pptx", b"broken", 422),
                                           ("empty.pdf", b"", 422)):
                response = self.upload(name, content)
                self.assertEqual(response.status_code, status, response.text)
            self.assertIn("Convert", self.upload("old.ppt", b"old").json()["detail"])
            self.assertEqual(self.upload("student.pdf", pdf_bytes(),
                                         {"Authorization": f"Bearer {self.join()['token']}"}).status_code, 403)
            self.assertEqual(self.upload("missing.pdf", pdf_bytes(), {}).status_code, 403)
            with patch.dict(os.environ, {"MAX_UPLOAD_MB": "1"}):
                self.assertEqual(self.upload("large.pdf", b"x" * (1024 * 1024 + 1)).status_code, 413)
            teacher.send_json({"type": "ping"})
            self.assertEqual(receive(teacher, lambda s: s.get("type") == "pong")["type"], "pong")
            teacher.send_json({"type": "set_slide", "index": 0})
            after = receive(teacher, lambda s: s.get("current_slide") == 0)
            self.assertEqual(after["presentation_id"], state["presentation_id"])
            self.assertEqual(after["active_material"]["filename"], "lecture.pptx")

    def test_restart_loses_session_not_local_file(self):
        self.assertEqual(self.upload("restart.pdf", pdf_bytes()).status_code, 201)
        fresh = TestClient(create_app())
        try:
            self.assertEqual(fresh.post(f"/api/sessions/{self.code}/join").status_code, 404)
        finally:
            fresh.close()

    def test_mocked_visual_pptx_upload_synchronises_two_students(self):
        def fake_conversion(command, **_kwargs):
            output = Path(command[command.index("--outdir") + 1])
            source = Path(command[-1])
            (output / f"{source.stem}.pdf").write_bytes(pdf_bytes())
            return SimpleNamespace(returncode=0)

        with ExitStack() as stack:
            teacher = self.connect(stack, self.lecturer)
            student_a_credentials = self.join()
            student_b_credentials = self.join()
            student_a = self.connect(stack, student_a_credentials)
            student_b = self.connect(stack, student_b_credentials)
            with patch("app.document_processing.libreoffice_executable", return_value=Path(sys.executable)), \
                 patch("app.document_processing.subprocess.run", side_effect=fake_conversion):
                result = self.upload("visual.pptx", pptx_bytes())
            self.assertEqual(result.status_code, 201, result.text)
            self.assertEqual(result.json()["rendering"], "visual")
            material_id = result.json()["presentation_id"]
            for socket in (teacher, student_a, student_b):
                state = receive(socket, lambda s: s.get("presentation_id") == material_id)
                self.assertEqual(len(state["slides"]), 2)
                self.assertIn("First PPTX slide", state["slides"][0]["text"])
                self.assertTrue(state["slides"][0]["image_url"])
            image_url = state["slides"][0]["image_url"]
            self.assertEqual(self.client.get(image_url, headers={"Authorization": f"Bearer {student_a_credentials['token']}"}).status_code, 200)
            teacher.send_json({"type": "set_slide", "index": 1})
            for socket in (teacher, student_a, student_b):
                self.assertEqual(receive(socket, lambda s: s.get("current_slide") == 1)["slides"][1]["title"], "Slide 2")

    def test_validation_rejects_spoofed_files_before_activation(self):
        self.assertEqual(self.upload("renamed.pdf", b"This is a TXT file.").status_code, 422)
        self.assertEqual(self.upload("unreadable.pptx", b"PK\x03\x04not-a-real-archive").status_code, 422)
        self.assertEqual(self.upload("unsupported.txt", pdf_bytes()).status_code, 415)
        damaged = self.upload("slightly-damaged.pdf", pdf_bytes()[:-5])
        self.assertEqual(damaged.status_code, 201, damaged.text)
        self.assertEqual(damaged.json()["slide_count"], 2)
        with patch.dict(os.environ, {"MAX_UPLOAD_MB": "25"}), patch("app.materials.process_pdf") as parser:
            oversized = self.upload("oversized.pdf", b"x" * (25 * 1024 * 1024 + 1))
        self.assertEqual(oversized.status_code, 413)
        parser.assert_not_called()
        self.assertEqual(self.upload("still-unreadable.pdf", b"not a PDF").status_code, 422)
        self.assertEqual(self.client.get("/api/health").status_code, 200)

    def test_unreadable_file_preserves_feedback_and_student_sync(self):
        with ExitStack() as stack:
            teacher = self.connect(stack, self.lecturer)
            credentials_a = self.join()
            credentials_b = self.join()
            a = self.connect(stack, credentials_a)
            b = self.connect(stack, credentials_b)
            good = self.upload("working.pdf", pdf_bytes())
            self.assertEqual(good.status_code, 201)
            presentation_id = good.json()["presentation_id"]
            receive(teacher, lambda s: s.get("presentation_id") == presentation_id)
            receive(a, lambda s: s.get("presentation_id") == presentation_id)
            receive(b, lambda s: s.get("presentation_id") == presentation_id)
            for socket, choice in ((a, "understand"), (b, "not_understand")):
                socket.send_json({"type": "submit_feedback", "slide_index": 0, "choice": choice,
                                  "presentation_id": presentation_id})
            before = receive(teacher, lambda s: s.get("current_feedback", {}).get("total") == 2)
            self.assertTrue(before["current_feedback"]["flagged"])
            unreadable = self.upload("unreadable.pdf", b"This cannot be opened by PyMuPDF.")
            self.assertEqual(unreadable.status_code, 422)
            self.assertIn("Could not process", unreadable.json()["detail"])
            session = self.client.app.state.store.session
            self.assertEqual(session.presentation_id, presentation_id)
            self.assertEqual(len(session.slides), 2)
            self.assertEqual(session.snapshot("lecturer", self.lecturer["token"])["current_feedback"], before["current_feedback"])
            teacher.send_json({"type": "set_slide", "index": 1})
            for socket in (teacher, a, b):
                state = receive(socket, lambda s: s.get("current_slide") == 1)
                self.assertEqual(state["presentation_id"], presentation_id)
            teacher.send_json({"type": "set_slide", "index": 0})
            restored = receive(teacher, lambda s: s.get("current_slide") == 0)
            self.assertEqual(restored["current_feedback"], before["current_feedback"])


if __name__ == "__main__":
    unittest.main()
