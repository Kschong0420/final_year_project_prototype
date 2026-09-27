"""Structured PPTX extraction and optional LibreOffice rendering tests."""
import io
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from app.document_processing import prepare_pptx, process_pptx, process_pdf, libreoffice_executable
from app.document_processing import _pdf_content


def detailed_pptx(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Structured lecture"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1.2), Inches(6), Inches(1.5))
    first = textbox.text_frame.paragraphs[0]
    first.text = "First bullet"
    nested = textbox.text_frame.add_paragraph()
    nested.text = "Nested bullet"
    nested.level = 1
    numbered = textbox.text_frame.add_paragraph()
    numbered.text = "First numbered"
    for paragraph in (first, nested):
        marker = OxmlElement("a:buChar")
        marker.set("char", "•")
        paragraph._p.get_or_add_pPr().append(marker)
    marker = OxmlElement("a:buAutoNum")
    marker.set("type", "arabicPeriod")
    numbered._p.get_or_add_pPr().append(marker)
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
    for row, values in enumerate((("Topic", "Value"), ("Cells", "Connected"))):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    image = Image.new("RGB", (4, 4), "red")
    image_stream = io.BytesIO()
    image.save(image_stream, format="PNG")
    image_stream.seek(0)
    slide.shapes.add_picture(image_stream, Inches(1), Inches(5), width=Inches(1))
    presentation.save(path)


def rendered_pdf():
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "LibreOffice visual test")
    content = document.tobytes()
    document.close()
    return content


def ole_table_pptx(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(.3), Inches(4), Inches(.5)).text = "Table slide"
    icon = Image.new("RGB", (400, 200), "white")
    drawing = ImageDraw.Draw(icon)
    drawing.rectangle((0, 0, 399, 199), outline="black", width=4)
    drawing.line((0, 75, 399, 75), fill="black", width=3)
    drawing.line((200, 0, 200, 199), fill="black", width=3)
    drawing.text((25, 25), "Topic", fill="black")
    drawing.text((225, 25), "Value", fill="black")
    drawing.text((25, 110), "Row A", fill="black")
    drawing.text((225, 110), "Pass", fill="black")
    icon_stream = io.BytesIO()
    icon.save(icon_stream, format="PNG")
    icon_stream.seek(0)
    slide.shapes.add_ole_object(io.BytesIO(b"embedded-sheet"), "Excel.Sheet.12",
                                Inches(2), Inches(1), Inches(4), Inches(2), icon_file=icon_stream)
    presentation.save(path)


class ProcessingTests(unittest.TestCase):
    def test_actual_adaptive_learning_pdf_keeps_table_and_discards_false_panel_table(self):
        """A one-page extract of the user's problematic PDF, verified against its screenshot."""
        fixture = Path(__file__).parent / "fixtures" / "adaptive-learning-page.pdf"
        with tempfile.TemporaryDirectory() as directory:
            slides = process_pdf(fixture, Path(directory) / "previews", "adaptive-fixture")
            slide = slides[0]
            lines = slide["text"].splitlines()
            self.assertEqual(slide["title"], "Page 1")
            self.assertEqual(lines[:5], [
                "Adaptive Learning Systems",
                "Study | Adaptation Approach",
                "Reddig et al. (2025) | Individualised feedback based on student errors",
                "Villegas-Ch et al. (2025) | Personalised adaptive learning",
                "Lee et al. (2026) | Just-in-time adaptive feedback",
            ])
            self.assertIn("There is also supporting research showing that real-time class-level information can help teachers make instructional decisions", lines)
            self.assertIn("For example, classroom communication systems have been used to collect students' answers and feedback the class results to trigger", lines)
            self.assertIn("further discussion or teaching action.", lines)
            self.assertNotIn("There i", lines)
            self.assertNotIn("For ex", lines)
            self.assertNotIn("furthe | There i", slide["text"])
            tables = [block for block in slide["structured_content"] if block["kind"] == "table"]
            self.assertEqual(len(tables), 1)
            self.assertEqual(tables[0]["rows"], [
                ["Study", "Adaptation Approach"],
                ["Reddig et al. (2025)", "Individualised feedback based on student errors"],
                ["Villegas-Ch et al. (2025)", "Personalised adaptive learning"],
                ["Lee et al. (2026)", "Just-in-time adaptive feedback"],
            ])
            self.assertTrue(all(block["slide_number"] == 1 for block in slide["structured_content"]))
            self.assertTrue((Path(directory) / "previews" / "slide-1.png").is_file())

    def test_numbered_pptx_sections_keep_nearby_heading_and_body(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sections.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            for value, x, y, width in (("Problem Statement", 1, .4, 5),
                                       ("First constraint", 2.2, 2.44, 5),
                                       ("01", 1.5, 2.45, .35),
                                       ("A complete supporting sentence.", 1.5, 3.1, 8),
                                       ("Second constraint", 2.2, 5.40, 5),
                                       ("02", 1.5, 5.41, .35),
                                       ("Another complete sentence.", 1.5, 6.1, 8)):
                slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(.45)).text = value
            presentation.save(path)
            extracted = process_pptx(path)[0]
            self.assertEqual(extracted["text"].splitlines(), [
                "Problem Statement", "01 First constraint", "A complete supporting sentence.",
                "02 Second constraint", "Another complete sentence.",
            ])
            self.assertTrue(all(block["slide_number"] == 1 for block in extracted["structured_content"]))

    def test_actual_problem_statement_pdf_keeps_numbered_sections_without_fake_tables(self):
        fixture = Path(__file__).parent / "fixtures" / "problem-statement-page.pdf"
        with tempfile.TemporaryDirectory() as directory:
            slide = process_pdf(fixture, Path(directory) / "previews", "problem-fixture")[0]
        lines = slide["text"].splitlines()
        self.assertEqual(lines[0:2], ["Problem Statement", "01 Lecturers Lack Time to Prepare Interactive Learning Materials"])
        self.assertIn("02 Students Experience Passive Learning During Lectures", lines)
        self.assertIn("03 Existing AI Learning Systems Do Not Support Live Classroom Adaptation", lines)
        self.assertIn("time for lesson preparation.", lines)
        self.assertIn("to traditional lectures.", lines)
        self.assertFalse(any(block["kind"] == "table" for block in slide["structured_content"]))
        self.assertNotIn(" | | ", slide["text"])

    def test_real_libreoffice_places_embedded_table_preview_at_source_bounds(self):
        if libreoffice_executable() is None:
            self.skipTest("LibreOffice is not installed")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "ole-table.pptx"
            ole_table_pptx(path)
            slides, mode, warning = prepare_pptx(path, root / "previews", "real-ole", root / "temp")
            self.assertEqual((mode, warning), ("visual", None))
            self.assertEqual(len(slides), 1)
            with Image.open(root / "previews" / "slide-1.png") as preview:
                # Source object is at x=2..6 in on a 10 in slide; no tiny image
                # should appear at the upper-left 100 pt square.
                x = round(preview.width * .4)
                y = round(preview.height * .27)
                center = preview.getpixel((x, y))
                self.assertLess(min(center), 40)  # black table grid inside source bounds
                upper_left = preview.getpixel((round(preview.width * .05), round(preview.height * .05)))
                self.assertGreater(min(upper_left), 230)

    def test_pdf_single_column_image_and_two_column_numbered_agenda(self):
        document = pymupdf.open()
        single = document.new_page(width=600, height=400)
        single.insert_text((40, 45), "Single column title")
        single.insert_text((40, 100), "First paragraph")
        single.insert_text((40, 145), "Second paragraph")
        icon = io.BytesIO()
        Image.new("RGB", (20, 20), "red").save(icon, format="PNG")
        single.insert_image(pymupdf.Rect(450, 80, 500, 130), stream=icon.getvalue())
        text, blocks = _pdf_content(single)
        self.assertEqual(text.splitlines(), ["Single column title", "First paragraph", "Second paragraph"])
        self.assertTrue(any(block["kind"] == "image" for block in blocks))

        agenda = document.new_page(width=600, height=400)
        agenda.insert_text((240, 50), "Agenda")
        for number, label, x, y in (("01", "Project Background", 40, 130),
                                     ("02", "Methodology", 40, 180),
                                     ("03", "Evaluation", 40, 230),
                                     ("04", "Limitations", 325, 130),
                                     ("05", "Conclusion", 325, 180),
                                     ("06", "Review", 325, 230)):
            agenda.insert_text((x, y), number)
            agenda.insert_text((x + 35, y), label)
        text, _ = _pdf_content(agenda)
        self.assertEqual(text.splitlines(), ["Agenda", "01 Project Background", "02 Methodology",
                                                 "03 Evaluation", "04 Limitations", "05 Conclusion", "06 Review"])
        document.close()

    def test_pdf_native_table_keeps_rows_and_columns(self):
        document = pymupdf.open()
        page = document.new_page(width=500, height=300)
        for x in (50, 200, 400):
            page.draw_line((x, 80), (x, 200))
        for y in (80, 120, 160, 200):
            page.draw_line((50, y), (400, y))
        for x, y, value in ((60, 105, "Topic"), (210, 105, "Value"),
                            (60, 145, "Alpha"), (210, 145, "One"),
                            (60, 185, "Beta"), (210, 185, "Two")):
            page.insert_text((x, y), value)
        text, blocks = _pdf_content(page)
        tables = [block for block in blocks if block["kind"] == "table"]
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0]["rows"], [["Topic", "Value"], ["Alpha", "One"], ["Beta", "Two"]])
        self.assertIn("Alpha | One", text)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "table.pdf"
            document.save(path)
            slides = process_pdf(path, root / "previews", "table-test")
            self.assertEqual(len(slides), 1)
            with Image.open(root / "previews" / "slide-1.png") as preview:
                self.assertEqual(preview.size, (625, 375))
                grid = preview.getpixel((round(50 * 1.25), round(80 * 1.25)))
                self.assertLess(min(grid[:3]), 40)
        document.close()

    def test_pptx_two_column_numbered_agenda(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "agenda.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            for value, x, y in (("Agenda", 4, .4), ("01", .7, 2), ("Background", 1.4, 2),
                                ("02", .7, 3), ("Methods", 1.4, 3), ("03", .7, 4),
                                ("Results", 1.4, 4), ("04", 5.7, 2), ("Limits", 6.4, 2),
                                ("05", 5.7, 3), ("Review", 6.4, 3), ("06", 5.7, 4),
                                ("Close", 6.4, 4)):
                box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(1.3 if value.isdigit() else 2.5), Inches(.4))
                box.text = value
            presentation.save(path)
            text = process_pptx(path)[0]["text"]
            self.assertEqual(text.splitlines(), ["Agenda", "01 Background", "02 Methods", "03 Results",
                                                  "04 Limits", "05 Review", "06 Close"])

    def test_structured_title_lists_table_and_image(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "detailed.pptx"
            detailed_pptx(path)
            slide = process_pptx(path)[0]
            self.assertEqual(slide["title"], "Structured lecture")
            blocks = slide["structured_content"]
            self.assertEqual([block["kind"] for block in blocks],
                             ["title", "paragraph", "paragraph", "paragraph", "table", "image"])
            self.assertEqual([block["list_type"] for block in blocks[1:4]],
                             ["bullet", "bullet", "numbered"])
            self.assertEqual([block["level"] for block in blocks[1:4]], [0, 1, 0])
            self.assertEqual(blocks[4]["rows"], [["Topic", "Value"], ["Cells", "Connected"]])
            self.assertIn("  • Nested bullet", slide["text"])
            self.assertIn("1. First numbered", slide["text"])
            self.assertIn("Topic | Value", slide["text"])

    def test_conversion_pipeline_uses_isolated_profile_and_source_text(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "detailed.pptx"
            detailed_pptx(path)
            observed = {}

            def fake_run(command, **kwargs):
                observed["command"] = command
                observed["timeout"] = kwargs["timeout"]
                output = Path(command[command.index("--outdir") + 1])
                (output / f"{path.stem}.pdf").write_bytes(rendered_pdf())
                return SimpleNamespace(returncode=0)

            with patch("app.document_processing.libreoffice_executable", return_value=Path(sys.executable)), \
                 patch("app.document_processing.subprocess.run", side_effect=fake_run):
                slides, mode, warning = prepare_pptx(path, root / "previews", "material123", root / "temporary")
            self.assertEqual(mode, "visual")
            self.assertIsNone(warning)
            self.assertEqual(len(slides), 1)
            self.assertTrue((root / "previews" / "slide-1.png").is_file())
            self.assertIn("Structured lecture", slides[0]["text"])
            self.assertEqual(slides[0]["structured_content"][4]["kind"], "table")
            self.assertIn("--headless", observed["command"])
            self.assertTrue(any(part.startswith("-env:UserInstallation=file:") for part in observed["command"]))
            self.assertEqual(observed["timeout"], 60)

    def test_missing_converter_and_timeout_use_text_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "detailed.pptx"
            detailed_pptx(path)
            with patch("app.document_processing.libreoffice_executable", return_value=None):
                slides, mode, warning = prepare_pptx(path, root / "preview", "m", root / "tmp")
            self.assertEqual(mode, "text-only")
            self.assertIn("not installed", warning)
            self.assertIsNone(slides[0]["image_url"])
            with patch("app.document_processing.libreoffice_executable", return_value=Path(sys.executable)), \
                 patch("app.document_processing.subprocess.run", side_effect=subprocess.TimeoutExpired("soffice", 60)):
                _, mode, warning = prepare_pptx(path, root / "preview", "m", root / "tmp")
            self.assertEqual(mode, "text-only")
            self.assertIn("timed out", warning)
            with patch.dict(os.environ, {"LIBREOFFICE_PATH": str(root / "missing.exe")}):
                _, mode, warning = prepare_pptx(path, root / "preview", "m", root / "tmp")
            self.assertEqual(mode, "text-only")
            self.assertIn("path is invalid", warning)


if __name__ == "__main__":
    unittest.main()
