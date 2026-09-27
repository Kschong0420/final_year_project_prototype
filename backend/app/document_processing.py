"""Extract source text independently from optional visual slide rendering."""
import os
import shutil
import subprocess
import tempfile
import io
from pathlib import Path

import pymupdf
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from .reading_order import order_items, pair_number_labels, text_lines


class ProcessingError(Exception):
    pass


class ConversionError(Exception):
    pass


def _pdf_content(page):
    blocks = []
    images = []
    for block in page.get_text("dict")["blocks"]:
        bbox = list(block["bbox"])
        if block["type"] == 1:
            images.append({"kind": "image", "bbox": bbox})
            continue
        line_items = [{"bbox": list(line["bbox"]),
                       "text": "".join(span["text"] for span in line["spans"]).strip()}
                      for line in block.get("lines", [])]
        line_items = order_items(pair_number_labels(line_items, page.rect.width),
                                 page.rect.width, page.rect.height)
        value = "\n".join(text_lines(line_items, page.rect.width))
        if value:
            blocks.append({"kind": "paragraph", "bbox": bbox, "text": value})
    # Native vector/text tables can be recovered as cells. Raster tables still
    # require OCR, which this milestone does not perform.
    try:
        # Filled panels can look like tables to the default detector, which
        # then manufactures clipped duplicate cells from complete paragraphs.
        found = page.find_tables(strategy="lines_strict")
        tables = found.tables
    except Exception:
        tables = []
    for table in tables:
        rows = [[cell.strip() if cell else "" for cell in row] for row in table.extract()]
        if len(rows) < 2 or max(map(len, rows), default=0) < 2 or sum(bool(cell) for row in rows for cell in row) < 2:
            continue
        x0, y0, x1, y1 = table.bbox
        blocks = [block for block in blocks if not (
            x0 <= (block["bbox"][0] + block["bbox"][2]) / 2 <= x1 and
            y0 <= (block["bbox"][1] + block["bbox"][3]) / 2 <= y1
        )]
        blocks.append({"kind": "table", "bbox": [x0, y0, x1, y1],
                       "rows": rows, "text": "\n".join(" | ".join(row) for row in rows)})
    ordered = order_items(pair_number_labels(blocks, page.rect.width),
                          page.rect.width, page.rect.height)
    return "\n".join(text_lines(ordered, page.rect.width)), ordered + images


def process_pdf(path: Path, preview_dir: Path, material_id: str, source_slides=None):
    """Render PDF pages; for PPTX retain python-pptx text instead of converted PDF text."""
    slides = []
    try:
        document_source = {"stream": path, "filetype": "pdf"} if isinstance(path, bytes) else {"filename": path}
        with pymupdf.open(**document_source) as document:
            if not document.is_pdf or document.page_count == 0:
                raise ProcessingError("The PDF has no usable pages.")
            if document.needs_pass:
                raise ProcessingError("This PDF is password-protected. Remove the password and try again.")
            if source_slides is not None and document.page_count != len(source_slides):
                raise ConversionError("LibreOffice produced a different number of pages than the PPTX slides.")
            preview_dir.mkdir(parents=True, exist_ok=True)
            for index, page in enumerate(document):
                image_name = f"slide-{index + 1}.png"
                page.get_pixmap(matrix=pymupdf.Matrix(1.25, 1.25), alpha=False).save(preview_dir / image_name)
                source = source_slides[index] if source_slides is not None else None
                extracted_text, structured = _pdf_content(page) if source is None else (source["text"], source["structured_content"])
                for block in structured:
                    block.setdefault("slide_number", index + 1)
                slides.append({
                    "title": source["title"] if source else f"Page {index + 1}",
                    "subtitle": "PowerPoint slide" if source else "PDF page",
                    "points": source["points"] if source else [],
                    "text": extracted_text,
                    "structured_content": structured,
                    "image_url": f"/api/materials/{material_id}/slides/{index + 1}/image",
                })
    except (ProcessingError, ConversionError):
        raise
    except Exception as exc:
        label = "converted PowerPoint slides" if source_slides is not None else "PDF"
        raise ProcessingError(f"Could not process this {label}. Check that the file is readable and not password-protected.") from exc
    return slides


def _paragraph_list_type(paragraph):
    properties = paragraph._p.pPr
    if properties is not None:
        if properties.find(qn("a:buNone")) is not None:
            return None
        if properties.find(qn("a:buAutoNum")) is not None:
            return "numbered"
        if properties.find(qn("a:buChar")) is not None or properties.find(qn("a:buBlip")) is not None:
            return "bullet"
    # A nested level commonly indicates a list in PowerPoint, but inherited
    # theme/layout bullet definitions are not always exposed by python-pptx.
    return "bullet" if paragraph.level > 0 else None


def _shape_blocks(shape, is_title=False):
    blocks = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in sorted(shape.shapes, key=lambda item: (item.top, item.left)):
            blocks.extend(_shape_blocks(child))
        return blocks
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            value = paragraph.text.strip()
            if value:
                blocks.append({
                    "kind": "title" if is_title else "paragraph",
                    "text": value,
                    "level": paragraph.level,
                    "list_type": None if is_title else _paragraph_list_type(paragraph),
                })
    if shape.has_table:
        rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
        blocks.append({"kind": "table", "rows": rows})
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        blocks.append({"kind": "image", "name": shape.name})
    return blocks


def _display_lines(blocks):
    lines = []
    counters = {}
    for block in blocks:
        if block["kind"] in ("title", "image"):
            continue
        if block["kind"] == "table":
            lines.extend(" | ".join(row) for row in block["rows"] if any(row))
            continue
        level = block["level"]
        if block["list_type"] == "numbered":
            counters[level] = counters.get(level, 0) + 1
            prefix = f"{counters[level]}. "
        elif block["list_type"] == "bullet":
            prefix = "• "
        else:
            prefix = ""
        for deeper in [key for key in counters if key > level]:
            del counters[deeper]
        lines.append("  " * level + prefix + block["text"])
    return lines


def process_pptx(path: Path):
    """Return ordered, structured source text, without depending on rendering."""
    try:
        presentation = Presentation(path)
        if not presentation.slides:
            raise ProcessingError("The PowerPoint file has no slides.")
        slides = []
        for index, slide in enumerate(presentation.slides):
            title_shape = slide.shapes.title
            title = title_shape.text.strip() if title_shape and title_shape.has_text_frame else ""
            items = []
            image_blocks = []
            for shape in slide.shapes:
                is_title = title_shape is not None and shape.shape_id == title_shape.shape_id
                blocks = _shape_blocks(shape, is_title)
                bbox = [shape.left, shape.top, shape.left + shape.width, shape.top + shape.height]
                for block in blocks:
                    block["bbox"] = bbox
                display = title if is_title else "\n".join(_display_lines(blocks))
                if display.strip():
                    items.append({"bbox": bbox, "text": display, "blocks": blocks, "is_title": is_title})
                else:
                    image_blocks.extend(blocks)
            ordered = order_items(pair_number_labels(items, presentation.slide_width),
                                  presentation.slide_width, presentation.slide_height)
            if not title and len(ordered) > 1:
                first, second = ordered[:2]
                if ("\n" not in first["text"] and len(first["text"]) <= 120 and
                        first["bbox"][3] < second["bbox"][1] - presentation.slide_height * .02):
                    title = first["text"]
                    first["is_title"] = True
                    for block in first["blocks"]:
                        if block["kind"] == "paragraph":
                            block["kind"] = "title"
            all_lines = text_lines(ordered, presentation.slide_width)
            points = text_lines([item for item in ordered if not item["is_title"]], presentation.slide_width)
            blocks = [block for item in ordered for block in item["blocks"]] + image_blocks
            for block in blocks:
                block["slide_number"] = index + 1
            slides.append({
                "title": title or f"Slide {index + 1}",
                "subtitle": "PowerPoint text view",
                "points": points,
                "text": "\n".join(all_lines),
                "structured_content": blocks,
                "image_url": None,
            })
        return slides
    except ProcessingError:
        raise
    except Exception as exc:
        raise ProcessingError("Could not process this PowerPoint file. Check that it is a valid .pptx file.") from exc


def libreoffice_executable():
    configured = os.getenv("LIBREOFFICE_PATH", "").strip()
    if configured:
        path = Path(configured)
        if not path.is_file():
            raise ConversionError("LibreOffice path is invalid. Check LIBREOFFICE_PATH; PowerPoint slides will use text view.")
        return path
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return Path(found)
    for path in (Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
                 Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe")):
        if path.is_file():
            return path
    return None


def convert_pptx_to_pdf(path: Path, temp_root: Path):
    executable = libreoffice_executable()
    if executable is None:
        raise ConversionError("LibreOffice is not installed. PowerPoint slides are shown as text only.")
    try:
        timeout = max(1, int(os.getenv("LIBREOFFICE_TIMEOUT_SECONDS", "60")))
    except ValueError:
        timeout = 60
    temp_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="pptx-", dir=temp_root) as working:
        directory = Path(working)
        output = directory / "output"
        output.mkdir()
        profile = directory / "profile"
        command = [str(executable), f"-env:UserInstallation={profile.as_uri()}",
                   "--headless", "--convert-to", "pdf:impress_pdf_Export", "--outdir", str(output), str(path)]
        try:
            result = subprocess.run(command, capture_output=True, text=True, timeout=timeout,
                                    check=False, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except subprocess.TimeoutExpired as exc:
            raise ConversionError(f"LibreOffice conversion timed out after {timeout} seconds. PowerPoint slides are shown as text only.") from exc
        except OSError as exc:
            raise ConversionError("Could not start LibreOffice. Check LIBREOFFICE_PATH; PowerPoint slides are shown as text only.") from exc
        converted = output / f"{path.stem}.pdf"
        if result.returncode != 0 or not converted.is_file() or not converted.stat().st_size:
            raise ConversionError("LibreOffice could not render this PowerPoint file. PowerPoint slides are shown as text only.")
        # The caller needs the PDF after TemporaryDirectory closes, so return its bytes.
        return converted.read_bytes()


def _ole_previews_for_conversion(path: Path, working_dir: Path):
    """Remove OLE shapes LibreOffice misplaces; retain their own PPTX preview images."""
    presentation = Presentation(path)
    previews = {}
    for index, slide in enumerate(presentation.slides):
        for shape in list(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT:
                continue
            blip = shape._element.find(".//" + qn("a:blip"))
            relation = blip.get(qn("r:embed")) if blip is not None else None
            if not relation:
                continue
            try:
                image_part = slide.part.related_part(relation)
                image_data = image_part.blob
                # Validate preview bytes before removing the source object.
                with Image.open(io.BytesIO(image_data)) as preview:
                    if preview.format not in ("PNG", "JPEG"):
                        continue
                    preview.verify()
            except Exception:
                continue  # Leave an object intact if its preview cannot be used.
            previews.setdefault(index, []).append((
                (shape.left, shape.top, shape.width, shape.height), image_data
            ))
            parent = shape._element.getparent()
            parent.remove(shape._element)
    if not previews:
        return path, previews, presentation.slide_width, presentation.slide_height
    render_path = working_dir / "slides-without-ole.pptx"
    presentation.save(render_path)
    return render_path, previews, presentation.slide_width, presentation.slide_height


def _overlay_ole_previews(pdf_bytes, previews, slide_width, slide_height):
    if not previews:
        return pdf_bytes
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as document:
        for index, objects in previews.items():
            if index >= document.page_count:
                raise ConversionError("The converted PDF has fewer pages than the PPTX slides.")
            page = document[index]
            scale_x = page.rect.width / slide_width
            scale_y = page.rect.height / slide_height
            for (left, top, width, height), image_data in objects:
                target = pymupdf.Rect(left * scale_x, top * scale_y,
                                      (left + width) * scale_x, (top + height) * scale_y)
                if target.is_empty:
                    continue
                # The source object's box supplies placement. Keep the preview's
                # aspect ratio so table text and graphics are not stretched.
                page.insert_image(target, stream=image_data, keep_proportion=True, overlay=True)
        return document.tobytes(garbage=4, deflate=True)


def prepare_pptx(path: Path, preview_dir: Path, material_id: str, temp_root: Path):
    slides = process_pptx(path)
    try:
        temp_root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="ole-", dir=temp_root) as working:
            render_path, previews, width, height = _ole_previews_for_conversion(path, Path(working))
            pdf_bytes = convert_pptx_to_pdf(render_path, temp_root)
            pdf_bytes = _overlay_ole_previews(pdf_bytes, previews, width, height)
        # PyMuPDF accepts PDF bytes directly; keep the converter's temp directory isolated.
        visual_slides = process_pdf(pdf_bytes, preview_dir, material_id, source_slides=slides)
        return visual_slides, "visual", None
    except (ConversionError, ProcessingError) as exc:
        return slides, "text-only", str(exc)
    except OSError:
        return slides, "text-only", "LibreOffice temporary storage is unavailable. PowerPoint slides are shown as text only."
